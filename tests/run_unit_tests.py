# -*- coding: utf-8 -*-
"""
Юнит-тесты маппинга в поля Битрикс24 и обработки ошибок Mistral API.
Работают офлайн: без API-ключа, без webhook, без сети (обращения к порталу
и к LLM подменяются заглушками).

Запуск из корня проекта:
    python tests/run_unit_tests.py

Код возврата 0 — все проверки пройдены. Прогоняйте после правок
pcb_normalizer.py, bitrix24.py, справочников и ретраев в utils.py.
"""
import os
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

# Логи и артефакты — во временный каталог, проект не засоряем.
os.chdir(tempfile.mkdtemp())
os.environ["LOG_FILE"] = os.path.join(os.getcwd(), "logs.log")
os.environ.pop("BITRIX24_WEBHOOK_URL", None)
os.environ.pop("BITRIX24_TOKEN", None)

import httpx  # noqa: E402

from app import bitrix24 as b  # noqa: E402
from app import bitrix24_api, pcb_normalizer as pn  # noqa: E402

FAILS = []


def check(label, actual, expected):
    ok = actual == expected
    print(f"  {'✓' if ok else '✗'} {label}: {actual!r}" + ("" if ok else f"   ОЖИДАЛОСЬ {expected!r}"))
    if not ok:
        FAILS.append(label)


def section(title):
    print(f"\n═══ {title} ═══")


# ─────────────────────────────────────────────────────────────────
section("1. Габариты: допуски не должны попадать в размеры")
for label, raw, exp in [
    ("квалитет h12 (бланк ЗАСЛОН)", "114 h12 x 47 h12", (114.0, 47.0)),
    ("сырые строки формы", "длина 114 Допуск h12, ширина 47 h12", (114.0, 47.0)),
    ("квалитет js13", "90 js13 x 60 js13", (90.0, 60.0)),
    ("± допуск", "303±0.2 x 111.0±0.2 mm", (303.0, 111.0)),
    ("+/- допуск раздельно", "160 x 120 +0,3 -0,3", (160.0, 120.0)),
    ("процентный допуск", "100 x 75 ±10%", (100.0, 75.0)),
    ("дефис как разделитель", "114-47", (114.0, 47.0)),
    ("запятая как дробная", "160мм x 120мм ±0,2мм", (160.0, 120.0)),
    ("круглая плата — не пара", "Ø120,0 ±0,1", None),
    ("пустое значение", "", None),
]:
    check(label, b._parse_dimensions(raw, 5.0), exp)

# ─────────────────────────────────────────────────────────────────
section("2. Да/Нет в русских и английских бланках")
for raw, exp in [
    ("есть", True), ("Да", True), ("требуется", True),
    ("требуется в формате ХХХ", True), ("Yes", True), ("имеется", True),
    ("нет", False), ("Нет", False), ("не требуется", False),
    ("отсутствует", False), ("No", False), ("—", False),
    ("", None), ("на усмотрение производителя", None),
    # развёрнутые ответы LLM: ключевое слово внутри фразы
    ("Yes, формат ХХХ (уникальный код изготовителя)", True),
    ("Yes, формат ХХ (римские) ХХ (арабские)", True),
    ("No, не предусмотрено", False),
    # «да» не должно находиться внутри других слов («дата», «датчик»);
    # «не указана» — это отсутствие данных, а не отрицание: поле не отправляем
    ("дата изготовления не указана", None),
    ("датчик", None),
]:
    check(f"is_affirmative({raw!r})", pn.is_affirmative(raw), exp)

# ─────────────────────────────────────────────────────────────────
section("3. Справочники Controlled Impedance / Serial Number")
IMP, SER, DATE = "ufCrm24_1707840096", "ufCrm24_1707851328", "ufCrm24_1707841090"
for raw, exp in [("есть", 7204), ("Yes", 7204), ("нет", 6088), ("не требуется", 6088)]:
    check(f"импеданс {raw!r}", pn.map_to_bitrix24_ids({"impedance_control": raw,
                                                       "_normalized": {}}).get(IMP), exp)
for raw, exp in [("требуется в формате ХХХ", 6518), ("Yes", 6518), ("нет", 6516)]:
    check(f"серийный номер {raw!r}", pn.map_to_bitrix24_ids({"serial_number": raw,
                                                             "_normalized": {}}).get(SER), exp)

# ─────────────────────────────────────────────────────────────────
section("4. Date code: сторона и способ нанесения")
for data, exp in [
    ({"marking_side": "TOP"}, 6194),
    ({"marking_side": "BOTTOM"}, 6196),
    ({"marking_side": "сверху"}, 6194),
    ({"marking_side": "TOP, шелкография"}, 6206),
    ({"marking_side": "BOTTOM", "date_code": "Yes, шелкография"}, 6208),
    ({"marking_side": "TOP", "date_code": "Yes, паяльной маской"}, 6198),
    ({"marking_side": "TOP", "date_code": "Yes, медью"}, 6202),
    ({"marking_side": "нет"}, None),
    ({}, None),
]:
    label = f"{data.get('marking_side', '—')!r}" + (f" + {data['date_code']!r}" if data.get("date_code") else "")
    check(label, pn.map_to_bitrix24_ids({**data, "_normalized": {}}).get(DATE), exp)

# ─────────────────────────────────────────────────────────────────
section("5. Панель: геометрия, Board Per Panel, Order unit = pnl")
zaslon = {
    "board_name": "ИВУА.687261.013", "board_thickness": "2 ±0,3", "layer_count": 14,
    "base_material": "FR-4 Tg 170", "coverage_type": "ENIG",
    "board_size": "114 x 47", "panel_size": "134 x 74", "boards_per_panel": "1",
    "technological_fields": "Yes", "impedance_control": "есть",
    "serial_number": "требуется в формате ХХХ", "marking_side": "TOP",
    "min_hole_size": "0.2",
}
f = b.map_pcb_to_bitrix24_fields(zaslon, mistral_client=None)
check("Board Length", f.get("ufCrm24_1708353384301"), 114.0)
check("Board Width (был баг h12→12)", f.get("ufCrm24_1708353402068"), 47.0)
check("Panel Length", f.get("ufCrm24_1708375852081"), 134.0)
check("Panel Width", f.get("ufCrm24_1708375871512"), 74.0)
check("Board Per Panel", f.get("ufCrm24_1708375915545"), 1)
check("Panel Usage %", f.get("ufCrm24_1708375925847"), 54.03)
check("Order unit = pnl", f.get("ufCrm24_1707838030"), 5258)
check("Production Unit = pnl", f.get("ufCrm24_1707849863"), 6272)
check("Controlled Impedance", f.get(IMP), 7204)
check("Serial Number", f.get(SER), 6518)
check("Date code = on Top", f.get(DATE), 6194)
check("No of Layers = 14", f.get("ufCrm24_1709815185"), 6798)

section("6. Признаки поставки в панелях")
for label, args, exp in [
    ("несколько плат в панели", ({}, None, 2), True),
    ("технологические поля", ({"technological_fields": "Yes"}, None, 1), True),
    ("задан размер панели", ({}, (194.0, 74.0), 1), True),
    ("«Панелизация: Нет»", ({"panelization": "Нет"}, None, None), False),
    ("панелизации нет вовсе", ({}, None, None), False),
    ("текст «поставить в панели»", ({"panelization": "платы поставить в панели"}, None, None), True),
]:
    check(label, b._is_panel_delivery(*args), exp)

# ─────────────────────────────────────────────────────────────────
section("7. Min. Hole size — числовое поле (значение, не item_id)")
os.environ["B24_FIELD_MIN_HOLE_SIZE"] = "ufCrm24_HOLE"
import importlib  # noqa: E402
importlib.reload(b)
for raw, exp in [("0.2", 0.2), ("0,25 мм", 0.25), ("0.3 mm", 0.3)]:
    out = {}
    b._apply_optional_fields(out, {"min_hole_size": raw})
    check(f"{raw!r} → число", out.get("ufCrm24_HOLE"), exp)
out = {}
b._apply_optional_fields(out, {"min_hole_size": "0.2"})
check("тип значения", type(out.get("ufCrm24_HOLE")).__name__, "float")
os.environ.pop("B24_FIELD_MIN_HOLE_SIZE")

# ─────────────────────────────────────────────────────────────────
section("8. Живые справочники Битрикс24 (портал подменён заглушкой)")
FIELDS = {
    "ufCrm24_1707840096": {"title": "Controlled Impedance", "type": "iblock_element",
                           "settings": {"IBLOCK_ID": "300"}},
    "ufCrm24_1707838248": {"title": "Materials", "type": "iblock_element",
                           "settings": {"IBLOCK_ID": "56"}},
    "ufCrm24_DIFF": {"title": "No. of Diff Boards", "type": "double"},
}
LISTS = {
    300: [{"ID": 6088, "NAME": "No"}, {"ID": 7204, "NAME": "Yes"}],
    # на портале появился материал, которого нет в статическом словаре
    56: [{"ID": 5774, "NAME": "FR4 TG-180"}, {"ID": 9999, "NAME": "FR4 TG-155"}],
}


def fake_post(self, url, json=None, **kw):
    method = url.rsplit("/", 1)[-1]
    if method == "crm.item.fields":
        body = {"result": {"fields": FIELDS}}
    elif method == "lists.element.get":
        body = {"result": LISTS.get(int(json.get("IBLOCK_ID")), [])}
    else:
        body = {"result": []}
    return httpx.Response(200, json=body, request=httpx.Request("POST", url))


original_post = httpx.Client.post
httpx.Client.post = fake_post
os.environ["BITRIX24_WEBHOOK_URL"] = "https://x.bitrix24.ru/rest/6/TOKEN/crm.item.add"
bitrix24_api.reset_bitrix24_api()

materials = pn.get_dict("materials")
check("новое значение с портала видно", materials.get("FR4 TG-155"), 9999)
check("статические значения сохранены", materials.get("FR4 TG-180"), 5774)
check("автопоиск числового поля", b.resolve_optional_field_code("no_of_diff_boards"), "ufCrm24_DIFF")

section("9. Сбой портала не отключает поле навсегда")
httpx.Client.post = lambda self, url, json=None, **kw: (_ for _ in ()).throw(
    httpx.ConnectError("сеть недоступна"))
bitrix24_api.reset_bitrix24_api()
importlib.reload(b)
check("при сбое кода нет", b.resolve_optional_field_code("no_of_diff_boards"), "")
check("неудача НЕ закэширована", "no_of_diff_boards" in b._autodetected_codes, False)
httpx.Client.post = fake_post
bitrix24_api.reset_bitrix24_api()
check("после восстановления находит", b.resolve_optional_field_code("no_of_diff_boards"), "ufCrm24_DIFF")

httpx.Client.post = original_post
os.environ.pop("BITRIX24_WEBHOOK_URL", None)
bitrix24_api.reset_bitrix24_api()

# ─────────────────────────────────────────────────────────────────
section("10. Обязательные поля и офлайн-режим")
importlib.reload(b)
minimal = b.map_pcb_to_bitrix24_fields(
    {"board_name": "T-1", "board_thickness": "1.6", "base_material": "FR4",
     "board_size": "100 x 50"}, mistral_client=None)
check("OEM PN", minimal.get("ufCrm24_1709799376061"), "T-1")
check("Board Thickness", minimal.get("ufCrm24_1708374728464"), 1.6)
check("Order unit по умолчанию ea", minimal.get("ufCrm24_1707838030"), 5256)
try:
    b.map_pcb_to_bitrix24_fields({"board_thickness": "1.6"}, mistral_client=None)
    check("без OEM PN — ошибка", "не выброшена", "ValueError")
except ValueError:
    check("без OEM PN — ошибка", "ValueError", "ValueError")

section("11. Токен webhook маскируется в логах")
check("токен скрыт", b._mask_webhook_url("https://x.bitrix24.ru/rest/6/s3cr3t/crm.item.add"),
      "https://x.bitrix24.ru/rest/6/***/crm.item.add")

# ─────────────────────────────────────────────────────────────────
section("12. Ошибки Mistral: тип по HTTP-статусу, отдельные повторы при 429")
import time as _time  # noqa: E402
from app import utils  # noqa: E402
from app.model import PCBCharacteristics  # noqa: E402

RATE_LIMIT_BODY = '{"object":"error","message":"Rate limit exceeded","type":"rate_limited","code":"1300"}'


def http_error(status, headers=None, body=RATE_LIMIT_BODY):
    """httpx.HTTPStatusError в том виде, в каком его бросает langchain_mistralai."""
    req = httpx.Request("POST", "https://api.mistral.ai/v1/chat/completions")
    resp = httpx.Response(status, headers=headers or {}, text=body, request=req)
    return httpx.HTTPStatusError(f"Error response {status} while fetching {req.url}: {body}",
                                 request=req, response=resp)


class FakeParser:
    """Заглушка LLM: по очереди бросает исключения / возвращает результаты."""

    def __init__(self, *outcomes):
        self.outcomes = list(outcomes)
        self.calls = 0

    def invoke(self, messages, **kwargs):
        self.calls += 1
        out = self.outcomes.pop(0)
        if isinstance(out, BaseException):
            raise out
        return out


def run_parser(parser, **kw):
    """(имя исключения | результат, список пауз) — паузы перехватываются, не спим."""
    sleeps = []
    real_sleep = _time.sleep
    _time.sleep = lambda s: sleeps.append(s)
    try:
        return utils.process_excel_pcb_with_retry("x", parser, **kw), sleeps
    except Exception as e:
        return e, sleeps
    finally:
        _time.sleep = real_sleep


# Классификация
check("429 → лимит запросов", utils.is_rate_limit_error(http_error(429)), True)
check("401 → не лимит", utils.is_rate_limit_error(http_error(401)), False)
check("401 → не сеть (сервер ответил)", utils.is_network_error(http_error(401)), False)
check("ConnectError → сеть", utils.is_network_error(httpx.ConnectError("boom")), True)
check("WinError 10054 в тексте → сеть",
      utils.is_network_error(Exception("[WinError 10054] Удаленный хост принудительно разорвал подключение")), True)
validation_err = ValueError("layer_count: Input should be int, input_value='ГОСТ Р 53429-2009'")
check("«53429» в ошибке валидации — не лимит", utils.is_rate_limit_error(validation_err), False)
check("«53429» в ошибке валидации — не сеть", utils.is_network_error(validation_err), False)
check("Retry-After в секундах", utils.retry_after_seconds(http_error(429, {"Retry-After": "7"})), 7.0)
check("Retry-After отсутствует", utils.retry_after_seconds(http_error(429)), None)
check("заголовки лимитов в логе",
      utils.rate_limit_headers(http_error(429, {"x-ratelimitbysize-remaining-minute": "0",
                                                "content-type": "application/json"})),
      {"x-ratelimitbysize-remaining-minute": "0"})

OK = PCBCharacteristics(board_name="T-1")

# 429 дважды, затем успех: паузы 5 → max(10, Retry-After 12)
result, sleeps = run_parser(FakeParser(http_error(429), http_error(429, {"Retry-After": "12"}), OK),
                            rate_limit_attempts=5, rate_limit_delay=5.0)
check("после двух 429 — успех", getattr(result, "get", lambda k: result)("board_name"), "T-1")
check("паузы 5 → max(10, Retry-After 12)", sleeps, [5.0, 12.0])

# 429 на всех попытках
parser = FakeParser(*[http_error(429)] * 3)
result, sleeps = run_parser(parser, rate_limit_attempts=3, rate_limit_delay=5.0)
check("429 на всех попытках → MistralRateLimitError", type(result).__name__, "MistralRateLimitError")
check("попыток ровно 3", parser.calls, 3)
check("паузы 5 → 10", sleeps, [5.0, 10.0])
check("HTTP-статус виден через причину", utils.http_status_of(result), 429)

# Сервер просит ждать дольше потолка — сдаёмся сразу, не спим
parser = FakeParser(http_error(429, {"Retry-After": "600"}), OK)
result, sleeps = run_parser(parser, rate_limit_attempts=5)
check("Retry-After 600 → сдаёмся сразу", type(result).__name__, "MistralRateLimitError")
check("без ожидания", sleeps, [])

# Ошибка валидации ответа модели: без повторов, пробрасывается как есть
parser = FakeParser(validation_err)
result, sleeps = run_parser(parser)
check("ошибка валидации → как есть, без повторов", type(result).__name__, "ValueError")
check("вызов один", parser.calls, 1)
check("пауз нет", sleeps, [])

# Сетевые сбои: прежняя короткая стратегия
parser = FakeParser(httpx.ConnectError("boom"), httpx.ConnectError("boom"))
result, sleeps = run_parser(parser, max_retries=2, delay=2.0)
check("сеть на всех попытках → MistralNetworkError", type(result).__name__, "MistralNetworkError")
check("сетевых попыток 2", parser.calls, 2)
check("сетевая пауза 2", sleeps, [2.0])

# 401: неповторяемая
parser = FakeParser(http_error(401, body='{"message":"Unauthorized"}'))
result, sleeps = run_parser(parser)
check("401 → без повторов", parser.calls, 1)
check("401 → статус", utils.http_status_of(result), 401)

# ─────────────────────────────────────────────────────────────────
section("13. Сообщение для пользователя: без двойной обёртки")
from app import interface as ui  # noqa: E402

rl = utils.MistralRateLimitError("Превышен лимит запросов Mistral API (HTTP 429): ...")
check("лимит: текст как есть", ui._friendly_error_message(rl), str(rl))
check("лимит: без префикса «Ошибка при обработке файла»",
      ui._friendly_error_message(rl).startswith("Ошибка при обработке файла"), False)
net = utils.MistralNetworkError("Ошибка сети при обращении к Mistral API: ...")
check("сеть: текст как есть", ui._friendly_error_message(net), str(net))
check("401 по статусу", "401 Unauthorized" in ui._friendly_error_message(http_error(401, body="x")), True)
check("429 без ретраев (не из пайплайна)",
      ui._friendly_error_message(http_error(429)), "Превышен лимит запросов Mistral API (HTTP 429). Попробуйте позже.")
check("«53429» в тексте — обычная ошибка",
      ui._friendly_error_message(validation_err).startswith("Ошибка при обработке файла"), True)
check("«1401» в тексте — не 401",
      "Unauthorized" in ui._friendly_error_message(ValueError("value 1401 is invalid")), False)

# ─────────────────────────────────────────────────────────────────
section("14. Промпт 2 (нормализация): 429 → короткий повтор, потом fallback")
from types import SimpleNamespace  # noqa: E402
from langchain_mistralai import ChatMistralAI  # noqa: E402

NORMALIZED = '{"finish_type": "Imm. gold (chem.Ni/Au)", "copper_thickness": null, ' \
             '"base_material": "FR4 TG-180", "pcb_type": null, "ipc_class": null}'
MESSAGES = [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]


def run_adapter(*outcomes):
    """Подменяет ChatMistralAI.invoke: сеть не нужна. → (ответ | исключение, паузы)."""
    queue = list(outcomes)
    calls = []
    sleeps = []
    real_invoke, real_sleep = ChatMistralAI.invoke, _time.sleep

    def fake_invoke(self, messages, *a, **kw):
        calls.append(1)
        out = queue.pop(0)
        if isinstance(out, BaseException):
            raise out
        return SimpleNamespace(content=out)

    ChatMistralAI.invoke = fake_invoke
    _time.sleep = lambda s: sleeps.append(s)
    os.environ["MISTRAL_API_KEY"] = "test-key"
    try:
        adapter = b._MistralChatAdapter(None)
        try:
            resp = adapter.chat("mistral-small-latest", MESSAGES, response_format={"type": "json_object"})
            return resp.choices[0].message.content, len(calls), sleeps
        except Exception as e:
            return e, len(calls), sleeps
    finally:
        ChatMistralAI.invoke, _time.sleep = real_invoke, real_sleep
        os.environ.pop("MISTRAL_API_KEY", None)


content, calls, sleeps = run_adapter(http_error(429), NORMALIZED)
check("429, затем ответ", content, NORMALIZED)
check("вызовов 2", calls, 2)
check("пауза 5", sleeps, [5.0])

result, calls, sleeps = run_adapter(*[http_error(429)] * 3)
check("3×429 → MistralRateLimitError (уйдёт в fallback)", type(result).__name__, "MistralRateLimitError")
check("вызовов 3", calls, 3)
check("паузы 5 → 10", sleeps, [5.0, 10.0])

result, calls, sleeps = run_adapter(http_error(429, {"Retry-After": "45"}), NORMALIZED)
check("Retry-After 45 > потолка 20 → сразу fallback", type(result).__name__, "MistralRateLimitError")
check("без ожидания", sleeps, [])

# Полный путь: нормализация упала → поля Битрикс24 всё равно сформированы
real_invoke = ChatMistralAI.invoke
ChatMistralAI.invoke = lambda self, m, *a, **kw: (_ for _ in ()).throw(http_error(429))
_time.sleep, real_sleep = (lambda s: None), _time.sleep
os.environ["MISTRAL_API_KEY"] = "test-key"
try:
    fields = b.map_pcb_to_bitrix24_fields(
        {"board_name": "T-2", "board_thickness": "1.6", "base_material": "FR4 TG-180",
         "coverage_type": "ENIG"}, mistral_client=b._MistralChatAdapter(None))
finally:
    ChatMistralAI.invoke, _time.sleep = real_invoke, real_sleep
    os.environ.pop("MISTRAL_API_KEY", None)
check("fallback: OEM PN", fields.get("ufCrm24_1709799376061"), "T-2")
check("fallback: материал из dicts.*, не MIX/Others", fields.get("ufCrm24_1707838248"), 5774)

# ─────────────────────────────────────────────────────────────────
section("15. Модели из окружения и лог расхода токенов")
from langchain_core.outputs import LLMResult  # noqa: E402

for var in ("MISTRAL_MODEL_PARSE", "MISTRAL_MODEL_NORMALIZE"):
    os.environ.pop(var, None)
check("Промпт 1 по умолчанию", utils.get_parse_model(), "mistral-medium-latest")
check("Промпт 2 по умолчанию", pn.get_normalize_model(), "mistral-small-latest")
os.environ["MISTRAL_MODEL_PARSE"] = "ministral-14b-2512"
os.environ["MISTRAL_MODEL_NORMALIZE"] = "ministral-8b-2512"
check("Промпт 1 из env", utils.get_parse_model(), "ministral-14b-2512")
check("Промпт 2 из env", pn.get_normalize_model(), "ministral-8b-2512")
# create_pcb_model: env → ChatMistralAI.model (сеть не нужна: без proxy пробы нет)
os.environ.pop("HTTPS_PROXY", None)
os.environ.pop("HTTP_PROXY", None)
parser_runnable = utils.create_pcb_model({"api_key": "test-key"})
check("create_pcb_model берёт модель из env", parser_runnable.first.bound.model, "ministral-14b-2512")
parser_runnable = utils.create_pcb_model({"api_key": "test-key", "model": "mistral-large-2512"})
check("params['model'] приоритетнее env", parser_runnable.first.bound.model, "mistral-large-2512")
for var in ("MISTRAL_MODEL_PARSE", "MISTRAL_MODEL_NORMALIZE"):
    os.environ.pop(var, None)

usage_logger = utils.TokenUsageLogger("t")
usage_logger.on_llm_end(LLMResult(generations=[[]], llm_output={
    "token_usage": {"prompt_tokens": 4300, "completion_tokens": 250, "total_tokens": 4550},
    "model_name": "mistral-large-2512"}))
check("usage из ответа", usage_logger.last,
      {"prompt_tokens": 4300, "completion_tokens": 250, "total_tokens": 4550})
usage_logger = utils.TokenUsageLogger("t")
usage_logger.on_llm_end(LLMResult(generations=[[]], llm_output=None))
check("ответ без usage — без падения", usage_logger.last, None)

# ─────────────────────────────────────────────────────────────────
section("16. Ограничения тарифа: нулевой лимит и 403 — без бесполезных повторов")
ZERO = {"x-ratelimit-limit-req-minute": "0", "x-ratelimit-remaining-req-minute": "0"}
check("нулевой лимит распознан", utils.rate_limit_is_zero(http_error(429, ZERO)), True)
check("обычный 429 — не нулевой", utils.rate_limit_is_zero(http_error(429, {"Retry-After": "3"})), False)
parser = FakeParser(http_error(429, ZERO), OK)
result, sleeps = run_parser(parser, rate_limit_attempts=5)
check("нулевой лимит → сразу MistralRateLimitError", type(result).__name__, "MistralRateLimitError")
check("без повторов и пауз", (parser.calls, sleeps), (1, []))
check("в сообщении — подсказка про модель/тариф", "MISTRAL_MODEL_PARSE" in str(result), True)
check("UI: текст как есть", ui._friendly_error_message(result), str(result))

tier = http_error(403, body='{"object":"error","message":"This model is not available in your subscription tier","type":"tier_not_allowed"}')
check("403 tier_not_allowed распознан", utils.is_tier_error(tier), True)
check("403 без tier — не тарифная", utils.is_tier_error(http_error(403, body='{"message":"Forbidden"}')), False)
parser = FakeParser(tier)
result, sleeps = run_parser(parser)
check("403 → без повторов", parser.calls, 1)
check("UI: сообщение про тариф", "тарифе" in ui._friendly_error_message(result), True)

# Промпт 2 упал → маппинг по справочникам + причина в diagnostics для статуса UI
real_invoke = ChatMistralAI.invoke
ChatMistralAI.invoke = lambda self, m, *a, **kw: (_ for _ in ()).throw(http_error(429, ZERO))
os.environ["MISTRAL_API_KEY"] = "test-key"
diag = {}
try:
    fields = b.map_pcb_to_bitrix24_fields(
        {"board_name": "T-3", "board_thickness": "1.6", "base_material": "FR4 TG-180"},
        mistral_client=b._MistralChatAdapter(None), diagnostics=diag)
finally:
    ChatMistralAI.invoke = real_invoke
    os.environ.pop("MISTRAL_API_KEY", None)
check("fallback: поля сформированы", fields.get("ufCrm24_1707838248"), 5774)
check("diagnostics: причина записана", "тарифе" in diag.get("normalization_error", ""), True)
fields2, err2, note2 = ui._map_fields_safely(
    {"board_name": "T-3", "board_thickness": "1.6", "base_material": "FR4 TG-180"}, None)
check("без LLM-клиента: заметки нет", (err2, note2), (None, None))

# ─────────────────────────────────────────────────────────────────
section("17. Файлы без текста и нечитаемые: понятная ошибка, в LLM не отправляются")
import gradio as gr  # noqa: E402
import openpyxl  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402

FILES = Path(tempfile.mkdtemp())


def read_error(path):
    """Текст DocumentReadError; None — файл прочитан; другое исключение — провал."""
    try:
        utils.extract_document_data(str(path))
        return None
    except utils.DocumentReadError as e:
        return str(e)
    except Exception as e:  # noqa: BLE001
        return f"ДРУГАЯ ОШИБКА {type(e).__name__}: {e}"


(FILES / "empty.txt").write_bytes(b"")
(FILES / "spaces.txt").write_text("  \n\t  ", encoding="utf-8")
DocxDocument().save(FILES / "empty.docx")
openpyxl.Workbook().save(FILES / "empty.xlsx")
(FILES / "broken.xlsx").write_bytes(os.urandom(2048))
(FILES / "broken.docx").write_bytes(os.urandom(2048))
(FILES / "html_as.xls").write_bytes(b"<html><body><table><tr><td>1</td></tr></table></body></html>")
(FILES / "spec.pdf").write_bytes(b"%PDF-1.4\n")
(FILES / "ok.txt").write_text("Толщина платы, мм | 1,6", encoding="utf-8")

for name in ("empty.txt", "spaces.txt", "empty.docx", "empty.xlsx"):
    msg = read_error(FILES / name) or ""
    check(f"{name}: «нет текста» с именем файла", "нет текста" in msg and name in msg, True)
for name, app in (("broken.xlsx", "Excel"), ("broken.docx", "Word"), ("html_as.xls", "Excel")):
    msg = read_error(FILES / name) or ""
    check(f"{name}: «не удалось прочитать», совет открыть в {app}",
          msg.startswith("Не удалось прочитать") and f"в {app}" in msg, True)
check("broken.xlsx: без технического текста pandas",
      "engine manually" in (read_error(FILES / "broken.xlsx") or ""), False)
check("broken.docx: без «Package not found»",
      "Package not found" in (read_error(FILES / "broken.docx") or ""), False)
check("spec.pdf: формат не поддерживается",
      "не поддерживается" in (read_error(FILES / "spec.pdf") or ""), True)
check("ok.txt: читается", read_error(FILES / "ok.txt"), None)
check("UI: текст ошибки файла как есть", ui._friendly_error_message(utils.DocumentReadError("X")), "X")

# Весь путь через обработчик кнопки «Распознать»: LLM подменён счётчиком вызовов
calls = []
llm_answer = {}
real_create, real_process = utils.create_pcb_model, utils.process_excel_pcb_with_retry


def fake_process(*args, **kwargs):
    calls.append("process")
    return {**PCBCharacteristics().model_dump(), **llm_answer}


def run_ui(path):
    """(сообщение gr.Error | None, результат обработчика, число обращений к LLM)."""
    calls.clear()
    try:
        return None, ui.parse_excel_pcb(str(path)), len(calls)
    except gr.Error as e:
        return e.message, None, len(calls)


utils.create_pcb_model = lambda params: calls.append("create") or object()
utils.process_excel_pcb_with_retry = fake_process
os.environ.pop("MISTRAL_API_KEY", None)  # нормализация не должна уйти в сеть
try:
    msg, _, n = run_ui(FILES / "empty.txt")
    check("пустой файл: ошибка, в LLM не отправлен", ("нет текста" in (msg or ""), n), (True, 0))
    msg, _, n = run_ui(FILES / "broken.xlsx")
    check("битый .xlsx: ошибка без префикса, в LLM не отправлен",
          ((msg or "").startswith("Не удалось прочитать"), n), (True, 0))

    (FILES / "letter.txt").write_text("Добрый день! Высылаю счёт за июль.", encoding="utf-8")
    llm_answer = {}
    msg, _, n = run_ui(FILES / "letter.txt")
    check("не спецификация: «не найдено характеристик»",
          ("не найдено характеристик" in (msg or ""), n), (True, 2))
    llm_answer = {"company_name": "ООО Ромашка", "technological_fields": "No", "impedance_control": "нет"}
    msg, _, n = run_ui(FILES / "letter.txt")
    check("только заказчик и «нет» — тоже не спецификация", "не найдено характеристик" in (msg or ""), True)
    llm_answer = {"board_name": "T-17", "board_thickness": "1.6", "base_material": "FR4",
                  "coverage_type": "ENIG", "layer_count": 4}
    msg, result, n = run_ui(FILES / "ok.txt")
    status = str(((result or [None] * 8)[7] or {}).get("value", ""))
    check("спецификация: без ошибки, поля Битрикс24 сформированы", (msg, status.startswith("✅")), (None, True))
finally:
    utils.create_pcb_model, utils.process_excel_pcb_with_retry = real_create, real_process

# ─────────────────────────────────────────────────────────────────
section("18. Расширение в верхнем регистре: «Бланк заказа.XLSX» принимается")
demo = ui.create_interface()
upload = [blk for blk in demo.blocks.values()
          if isinstance(blk, gr.File) and blk.file_types and ".xlsx" in blk.file_types]
check("компонент загрузки найден", len(upload), 1)
types = upload[0].file_types if upload else []
for name in ("Бланк заказа ПП САНТ.758726.291.XLSX", "spec.XLS", "ЛТТ.DOCX", "old.DOC", "note.TXT",
             "spec.xlsx", "ЛТТ.docx"):
    # та же проверка, что в браузерной части Gradio: "." + всё после последней точки
    check(f"интерфейс принимает {name}", "." + name.split(".")[-1] in types, True)
check("неподдерживаемый .pdf по-прежнему не принимается", ".pdf" in types or ".PDF" in types, False)

(FILES / "ok.TXT").write_text("Толщина платы, мм | 1,6", encoding="utf-8")
wb = openpyxl.Workbook()
wb.active["A1"], wb.active["B1"] = "Толщина платы, мм", "1,6"
wb.save(FILES / "form.XLSX")
check("extract_document_data: .TXT", "1,6" in utils.extract_document_data(str(FILES / "ok.TXT")), True)
check("extract_document_data: .XLSX", "1,6" in utils.extract_document_data(str(FILES / "form.XLSX")), True)

# ─────────────────────────────────────────────────────────────────
print("\n" + "=" * 64)
if FAILS:
    print(f"ПРОВАЛЕНО: {len(FAILS)}")
    for name in FAILS:
        print("  ✗", name)
else:
    print("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ")
sys.exit(1 if FAILS else 0)
